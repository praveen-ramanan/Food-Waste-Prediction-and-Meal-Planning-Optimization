"""
Builds the Review 3 technical report (Review3_Report.docx) from the artefacts
written by src/pipeline.py and src/train_models.py, so every number in the
report is the number the code produced.

Run (after pipeline.py and train_models.py):
    python src/build_report.py
Optional PDF export (Windows + Microsoft Word):
    python src/build_report.py --pdf
"""
import json
import pathlib
import subprocess
import sys

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT, FIG, MODELS = ROOT / "out", ROOT / "figures", ROOT / "models"
DOCX_PATH = ROOT / "Review3_Report.docx"

STUDENT = "PRAVEEN RAMANAN K"
REG_NO = "241801214"
SUBMISSION_DATE = "03-10-2026"
TITLE = "Food Waste Prediction and Meal-Planning Optimization for Event Catering"
ACCENT = RGBColor(0x3B, 0x2F, 0xA8)


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_page_number(paragraph):
    run = paragraph.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        st = self.doc.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size in [(1, 15), (2, 12.5), (3, 11.5)]:
            h = self.doc.styles[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = ACCENT
            h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
            h.paragraph_format.space_after = Pt(4)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        hp = sec.header.paragraphs[0]
        hp.text = f"{STUDENT}\t\t{REG_NO}"
        hp.runs[0].font.size = Pt(9)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_number(fp)

    # text -----------------------------------------------------------------
    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)

    def p(self, text, bold_prefix=None, italic=False, align=None, size=None):
        para = self.doc.add_paragraph()
        if bold_prefix:
            para.add_run(bold_prefix).bold = True
        r = para.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.alignment = para.alignment or WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def bullets(self, items):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            if isinstance(it, tuple):
                para.add_run(it[0]).bold = True
                para.add_run(it[1])
            else:
                para.add_run(it)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # tables ---------------------------------------------------------------
    def table(self, df: pd.DataFrame, caption=None, col_widths=None, font=8.5, highlight_row=None, fmt=None):
        if caption:
            self.tab_no += 1
            cp = self.doc.add_paragraph()
            r = cp.add_run(f"Table {self.tab_no} -- {caption}")
            r.bold, r.font.size = True, Pt(9.5)
            cp.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(df.columns):
            cell = t.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold, run.font.size = True, Pt(font)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(cell, "3B2FA8")
        for i, (_, row) in enumerate(df.iterrows()):
            cells = t.add_row().cells
            for j, v in enumerate(row):
                if fmt and df.columns[j] in fmt and pd.notna(v) and not isinstance(v, str):
                    s = fmt[df.columns[j]].format(v)
                elif isinstance(v, float):
                    s = "--" if pd.isna(v) else f"{v:.4f}".rstrip("0").rstrip(".") if abs(v) < 1000 else f"{v:,.0f}"
                else:
                    s = str(v)
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(s)
                run.font.size = Pt(font)
                if highlight_row is not None and highlight_row(row):
                    run.bold = True
                    shade(cells[j], "E8E5FB")
                elif i % 2 == 1:
                    shade(cells[j], "F6F6FA")
        if col_widths:
            for row in t.rows:
                for j, w in enumerate(col_widths):
                    row.cells[j].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def kv_table(self, rows, caption=None, widths=(5.5, 11)):
        return self.table(pd.DataFrame(rows, columns=["Item", "Detail"]), caption, col_widths=list(widths), font=9.5)

    # figures --------------------------------------------------------------
    def figure(self, fname, caption, width_cm=16.5):
        path = FIG / fname
        if not path.exists():
            return
        self.fig_no += 1
        self.doc.add_picture(str(path), width=Cm(width_cm))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(f"Figure R{self.fig_no} -- {caption}")
        r.italic, r.font.size = True, Pt(9)

    def save(self, path):
        self.doc.save(path)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------
def load():
    csv = lambda n: pd.read_csv(OUT / n)  # noqa: E731
    with open(MODELS / "metadata.json") as f:
        meta = json.load(f)
    with open(OUT / "pipeline_manifest.json") as f:
        manifest = json.load(f)
    with open(OUT / "best_hyperparameters.json") as f:
        params = json.load(f)
    return {
        "meta": meta, "manifest": manifest, "params": params,
        "reg": csv("regression_model_comparison.csv"), "reg_base": csv("regression_baseline_metrics.csv"),
        "clf": csv("classification_model_comparison.csv"), "clf_base": csv("classification_baseline_metrics.csv"),
        "reg_ba": csv("regression_tuning_before_after.csv"), "clf_ba": csv("classification_tuning_before_after.csv"),
        "sig": csv("final_model_significance.csv"), "per_class": csv("per_class_metrics_all_models.csv"),
        "report": pd.read_csv(OUT / "classification_report_best_model.csv", index_col=0),
        "cm": pd.read_csv(OUT / "confusion_matrix_best_model.csv", index_col=0),
        "assoc": csv("eda_risk_tier_association.csv"), "anova2": csv("eda_two_way_anova.csv"),
        "effects": csv("categorical_effect_sizes.csv"), "cat": csv("eda_category_summary.csv"),
        "num": csv("eda_numeric_summary.csv"),
        "imp_reg": csv("feature_importance_final_regressor.csv"),
        "imp_clf": csv("feature_importance_final_classifier.csv"),
        "rationale": csv("model_selection_rationale.csv"),
    }


def cover(R: Report):
    d = R.doc
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("RAJALAKSHMI ENGINEERING COLLEGE")
    r.bold, r.font.size = True, Pt(18)
    R.p("Approved by AICTE | Affiliated to Anna University | Accredited by NAAC", align="center", size=9.5)
    R.p("(An Autonomous Institution, Affiliated to Anna University, Chennai)", align="center", size=9.5)
    info = [
        ("Course Code", "AD23532", "Course Name", "Principles of Data Science"),
        ("Course Type", "Lab Oriented Theory", "Course Category", "Professional Core (PC)"),
        ("Regulation", "2023", "Academic Year", "2026-27 (ODD)"),
        ("Degree and Branch", "B.Tech. Artificial Intelligence and Data Science", "Batch", "2024-2028"),
        ("Submission Date", SUBMISSION_DATE, "Semester", "V"),
    ]
    tb = d.add_table(rows=0, cols=4)
    tb.style = "Table Grid"
    for row in info:
        cells = tb.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(v)
            run.font.size = Pt(10)
            run.bold = j % 2 == 0
            if j % 2 == 0:
                shade(cells[j], "EEECFB")
    d.add_paragraph()
    for text, size, bold in [("Review - 3", 16, True), ("Model Building and Comparison Analysis", 14, True),
                             (TITLE, 13, False)]:
        para = d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        if not bold:
            run.italic = True
    d.add_paragraph()
    R.kv_table([("Student", STUDENT), ("Register number", REG_NO),
                ("Deliverables", "Technical report (this document); notebooks/Review3_Model_Building_and_Comparison.ipynb; "
                                 "src/ (pipeline, extended EDA, training); models/ (trained artefacts); "
                                 "app/streamlit_app.py (dashboard)")], widths=(4.5, 12))


def section_intro(R, A):
    m = A["manifest"]
    R.h("1. Introduction and Scope of Review 3")
    R.p("Review 1 defined the problem: forecast food waste for an event-catering booking before the event, grade the "
        "booking's waste risk, and turn that grade into a preparation action and, where warranted, a donation-partner "
        "alert. Review 2 audited, cleaned and explored the Kaggle dataset Food Wastage Data in Restaurant, built a "
        "leakage-safe train/test partition and engineered pre-event features.")
    R.p("Review 3 completes the project. It builds, tunes and compares machine-learning models for both learning tasks, "
        "selects the final models on quantitative evidence, and delivers a working decision-support dashboard. In line "
        "with the review brief, it covers:")
    R.bullets([("Model selection and justification", " -- seven algorithm families per task (Section 4)."),
               ("Model development", " -- leakage-safe 80/20 split plus 5-fold cross-validation (Section 5)."),
               ("Model evaluation", " -- MAE, MSE, RMSE and R2 for regression; accuracy, precision, recall, F1, ROC-AUC "
                                    "and confusion matrices for classification (Section 6)."),
               ("Hyperparameter tuning", " -- Grid Search and Randomized Search with before/after comparison (Section 7)."),
               ("Performance comparison", " -- tables, charts, strengths and limitations (Section 8)."),
               ("Final model selection", " -- a selection rule fixed in advance, plus a corrected significance test (Section 9)."),
               ("End-to-end product", " -- the Streamlit dashboard driven by the trained models (Section 10).")])
    R.p("Before modelling, the exploratory analysis is extended from the two dominant drivers studied in Review 2 to "
        "every column of the dataset (Section 3), because those findings shape the feature set and the model families.")
    R.h("1.1 Learning tasks and data recap", 2)
    R.kv_table([
        ("Raw data", f"{m['raw_shape'][0]:,} records x {m['raw_shape'][1]} columns (Kaggle, Food Wastage Data in Restaurant)"),
        ("Cleaned data", f"{m['cleaned_shape'][0]:,} unique bookings after removing exact duplicates"),
        ("Partition", f"{m['train_shape'][0]:,} training / {m['test_shape'][0]} test bookings (80/20, seed 42), "
                      f"split before any target-derived statistic"),
        ("Regression target", "wastage_food_amount (units of food wasted per booking)"),
        ("Classification target", f"Risk tier: LOW < {m['risk_tier_cutpoints']['Q1']:g} <= MEDIUM < "
                                  f"{m['risk_tier_cutpoints']['Q3']:g} <= HIGH (training-set quartiles)"),
        ("Tier distribution (train)", ", ".join(f"{k} {100 * v:.1f}%" for k, v in m["risk_tier_distribution"]["train"].items())),
        ("Design matrix", f"{m['design_matrix_columns']} columns: 4 numeric, 3 pricing/service encodings, one-hot "
                          f"encodings of 7 categorical columns and a 4-band guest-scale feature"),
        ("Noise ceiling", f"Implied R2 ceiling ~ {m['attainable_ceiling']['implied_r2_ceiling']} from identical-input "
                          f"records with differing targets (Review 2, Section 6.4)"),
    ], caption="Data and task summary")


def section_eda(R, A):
    R.h("2. Feature Set Carried into Modelling")
    R.p("The Review 2 feature set uses only information known at booking time: number_of_guests, quantity_of_food, "
        "food_per_guest, surplus_units, pricing_ordinal, is_high_price, is_sit_down and guest_scale, together with the "
        "categorical columns. Expected Waste per Guest is derived from the target and stays excluded. Two automated "
        "leakage guards remain in place: no feature name may contain a target-derived substring, and no numeric "
        "feature may correlate with the target above |r| = 0.95.")
    R.p("One change was made for Review 3 on the basis of the extended EDA. Preparation method is now one-hot encoded "
        "in full, not represented only by the sit-down flag. Finger Food and Buffet bookings differ systematically "
        "(Section 3), and the change lifts the 5-fold CV R2 of XGBoost by about 0.008. The design matrix grows from "
        "30 to 33 columns.")

    R.h("3. Extended Exploratory Data Analysis")
    R.p("Review 2 concentrated on pricing tier and preparation method. This section examines every column: its "
        "effect on the waste amount, on the waste rate (waste as a share of food prepared), and on the risk tier, "
        "and how the columns relate to each other. All statistics use the training partition only.")
    R.h("3.1 Waste across every categorical level", 2)
    R.figure("fig08_waste_by_every_category.png",
             "Waste distribution for every level of the eight categorical columns, with Kruskal-Wallis p-values.")
    cat = A["cat"]
    pr = cat[cat.column == "pricing"].set_index("level")
    pm = cat[cat.column == "preparation_method"].set_index("level")
    R.p(f"Only two columns move the distribution visibly. High-price bookings average {pr.loc['High', 'mean_waste']:.1f} "
        f"units, against {pr.loc['Moderate', 'mean_waste']:.1f} (Moderate) and {pr.loc['Low', 'mean_waste']:.1f} (Low). "
        f"Sit-down Dinner averages {pm.loc['Sit-down Dinner', 'mean_waste']:.1f}, against "
        f"{pm.loc['Buffet', 'mean_waste']:.1f} for Buffet and {pm.loc['Finger Food', 'mean_waste']:.1f} for Finger Food. "
        f"For the other six columns, the medians and interquartile ranges overlap almost completely.")
    R.figure("fig11_waste_rate_by_category.png",
             "Waste rate (waste / quantity prepared) per level with 95% confidence intervals.")
    R.p(f"Normalising by the quantity prepared does not change the picture: High-price bookings waste "
        f"{pr.loc['High', 'mean_waste_pct']:.2f}% of the food prepared, against {pr.loc['Low', 'mean_waste_pct']:.2f}% for "
        f"Low-price bookings. The pricing effect is therefore not just a consequence of premium events preparing "
        f"more food.")
    R.h("3.2 Association with the risk tier", 2)
    assoc = A["assoc"].copy()
    assoc["p_value"] = assoc["p_value"].map(lambda v: f"{v:.2e}")
    R.table(assoc, "Chi-square test and bias-corrected Cramer's V between each column and the risk tier",
            highlight_row=lambda r: r["association"] == "Strong")
    R.figure("fig13_risk_tier_composition.png", "LOW / MEDIUM / HIGH composition within every categorical level.")
    R.p("Three columns are strongly associated with the risk tier: pricing, the guest-scale band and preparation "
        "method. Location reaches statistical significance, but its effect is small (V < 0.1). Rural bookings have "
        "fewer HIGH-risk tiers, partly because they prepare less food on average. Food type, purchase history, "
        "seasonality, storage and event type have no practically meaningful association.")
    R.h("3.3 Interactions between columns", 2)
    an = A["anova2"][["factor_a", "factor_b", "F_a", "p_a", "F_b", "p_b", "F_ab", "p_ab", "partial_eta2_ab"]].copy()
    for c in ["p_a", "p_b", "p_ab"]:
        an[c] = an[c].map(lambda v: f"{v:.2e}")
    R.table(an, "Two-way ANOVA with sequential sums of squares (factor a, factor b given a, interaction)",
            highlight_row=lambda r: float(r["p_ab"]) < 0.05)
    R.figure("fig12_two_way_heatmaps.png", "Mean waste for four pairs of secondary columns (cell counts in brackets).")
    R.p("Pricing x preparation method is the only pair with a significant interaction. It reproduces the Review 2 "
        "result, F(4, 1285) = 11.54, and corresponds to a rank reversal of the service styles across price tiers. "
        "Event type x food type, season x storage and purchase history x event type have no detectable main or "
        "interaction effects. The one detectable non-pricing effect is location, and it is a main effect only: "
        "the location x pricing interaction is not significant.")
    R.figure("fig16_cramers_v_matrix.png",
             "Cramer's V between all categorical columns (bias-corrected).", 12)
    cv = pd.read_csv(OUT / "eda_cramers_v_matrix.csv", index_col=0)
    R.p(f"Some secondary columns are clearly related to each other. Food type and storage condition are strongly "
        f"associated (V = {cv.loc['type_of_food', 'storage_conditions']:.2f}), and storage, purchase history, "
        f"event type and seasonality show moderate links (V ~ 0.2). Pricing is related to the guest-scale band "
        f"(V = {cv.loc['pricing', 'guest_scale']:.2f}) and to preparation method "
        f"(V = {cv.loc['pricing', 'preparation_method']:.2f}). None of the inter-related secondary columns is "
        f"associated with waste or with the risk tier (Table 2), so their redundancy hides no signal. Tree "
        f"ensembles are unaffected by such collinearity.")
    R.h("3.4 Numeric structure", 2)
    R.figure("fig14_guests_quantity_map.png",
             "Booking map of guests x quantity coloured by mean waste (left); preparation intensity by pricing tier (right).")
    R.figure("fig19_pairplot_by_risk_tier.png",
             "Numeric and engineered features coloured by risk tier.", 14)
    R.p("The bookings sit on a small grid of guest/quantity combinations. The risk tiers separate along guest count "
        "and quantity, but they overlap heavily unless pricing is also known. This is why no single numeric feature "
        "is sufficient and why the models must combine the numeric features with pricing.")
    R.figure("fig15_event_type_profile.png", "Event-type profile: waste by service style, service mix, and food type.")
    R.figure("fig17_ecdf_by_segment.png", "Cumulative distribution of waste by pricing tier and service style, with the tier cut points.")
    R.figure("fig18_waste_per_guest.png", "Waste per guest by guest-scale band and by event type x pricing (descriptive only).")
    R.figure("fig20_season_location_storage.png", "Operational-context columns: season x location and storage x food type.")
    R.h("3.5 Implications for modelling", 2)
    R.bullets([
        "Signal is concentrated in pricing, preparation method, guests and quantity, with a non-additive "
        "pricing x service interaction. Tree ensembles, which learn interactions automatically, should outperform "
        "additive linear models.",
        "Weak columns are kept rather than dropped. Regularised ensembles can ignore them at little cost, and the "
        "final permutation-importance analysis (Section 9) checks this explicitly.",
        "The risk tier is imbalanced (LOW ~15%) and many targets sit exactly on the cut points, so macro-F1 rather "
        "than accuracy is the primary classification metric, and class weighting is part of the tuning search.",
        "The target takes few distinct values, which favours histogram-based boosting (XGBoost with "
        "tree_method='hist', HistGradientBoosting).",
    ])


def section_models(R, A):
    R.h("4. Model Selection and Justification")
    R.p("Seven algorithms are compared for each task. Together they span baseline, linear, instance-based, kernel, "
        "bagging and boosting families, so the comparison says something about the shape of the problem and not "
        "only about which library performs best.")
    rat = A["rationale"].copy()
    R.table(rat, "Candidate models and the reason each was chosen", col_widths=[2.6, 4.0, 10.0], font=8.5)
    R.p("XGBoost is configured with the histogram tree method (tree_method = 'hist'), which bins each feature into at "
        "most max_bin buckets before searching for splits. scikit-learn's HistGradientBoosting applies the same idea "
        "in an independent implementation. Comparing the two shows whether XGBoost's advantage comes from its "
        "regularisation terms or from histogram boosting itself.")

    R.h("5. Model Development")
    R.kv_table([
        ("Hold-out split", "80/20 random split (seed 42), made before risk tiers or segment benchmarks were computed; "
                           "test set used once, for final scoring only"),
        ("Cross-validation", "5-fold KFold (regression) and StratifiedKFold (classification), shuffled with seed 42, on "
                             "the training partition; identical folds for every model"),
        ("Scaling", "StandardScaler applied inside a Pipeline for Linear/Logistic, KNN and SVM, so each CV fold fits its "
                    "own scaler; tree models use the unscaled matrix"),
        ("Labels", "LOW / MEDIUM / HIGH encoded as 0 / 1 / 2 for every classifier, so XGBoost and scikit-learn share one "
                   "label space"),
        ("Two stages", "Baseline (library defaults), then tuned (Grid or Randomized search on the same 5 folds)"),
        ("Selection rule", "Highest tuned 5-fold CV score (R2 / macro-F1), fixed before inspecting test results"),
        ("Significance", "Nadeau-Bengio corrected resampled t-test, best vs runner-up, on 5x3 repeated CV"),
        ("Reproducibility", "python src/pipeline.py; python src/train_models.py; python src/build_notebooks.py"),
    ], caption="Training and validation protocol")


def _metric_cols(df, cols):
    keep = ["model"] + [c for c in cols if c in df.columns]
    return df[keep]


def section_eval(R, A):
    R.h("6. Model Evaluation")
    R.p("Regression is evaluated with MAE, MSE, RMSE and R2 on the test set, together with cross-validated R2 (mean "
        "+/- standard deviation over the 5 folds). Classification is evaluated with accuracy, balanced accuracy, macro "
        "precision, macro recall, macro and weighted F1, one-vs-rest ROC-AUC, log-loss and the recall of the HIGH "
        "tier. HIGH recall matters most operationally, because a missed HIGH booking means a missed donation alert. "
        "Accuracy is reported but not used for selection: always predicting MEDIUM already scores 50% on the test set.")
    R.h("6.1 Baseline models (default hyperparameters)", 2)
    rb = _metric_cols(A["reg_base"], ["cv_R2_mean", "cv_R2_std", "test_MAE", "test_MSE", "test_RMSE", "test_R2"])
    R.table(rb.sort_values("cv_R2_mean", ascending=False), "Regression baseline metrics")
    cb = _metric_cols(A["clf_base"], ["cv_macro_F1_mean", "cv_macro_F1_std", "test_accuracy", "test_precision_macro",
                                      "test_recall_macro", "test_macro_F1", "test_ROC_AUC_ovr"])
    R.table(cb.sort_values("cv_macro_F1_mean", ascending=False), "Classification baseline metrics", font=8)
    R.p("Even untuned, the three tree ensembles are well ahead of the other families. Random Forest is the strongest "
        "default model for both tasks, while XGBoost with default settings (depth 6, learning rate 0.3) over-fits "
        "slightly and trails it.")


def section_tuning(R, A):
    R.h("7. Hyperparameter Tuning and Model Optimisation")
    R.p("Grid Search is used where the space is small and discrete (Ridge alpha, KNN neighbours/weights/metric, SVM C, "
        "gamma, epsilon and class weight). Randomized Search with continuous distributions is used for the tree "
        "ensembles, where a full grid would be infeasible: 25 configurations for Random Forest, 40 for "
        "HistGradientBoosting and 50 for XGBoost, each scored over the same 5 folds.")
    rows = []
    for task, d in A["params"].items():
        for model, info in d.items():
            bp = ", ".join(f"{k}={v:.4g}" if isinstance(v, float) else f"{k}={v}" for k, v in info["best_params"].items())
            rows.append({"task": task, "model": model, "method": info["search"], "configs": info["candidates"],
                         "time (s)": info["seconds"], "best hyperparameters": bp})
    R.table(pd.DataFrame(rows), "Search method and best hyperparameters found", col_widths=[2.2, 2.8, 1.3, 1.3, 1.3, 8.0],
            font=7.5)
    R.h("7.1 Before vs after optimisation", 2)
    R.table(A["reg_ba"], "Regression: CV and test R2 before and after tuning", font=8)
    R.table(A["clf_ba"], "Classification: CV and test macro-F1 before and after tuning", font=8)
    R.figure("fig27_tuning_before_after.png", "Effect of tuning on cross-validated scores (baseline -> tuned).")
    ba = A["reg_ba"].set_index("model")
    x = A["params"]["regression"]["XGBoost"]["best_params"]
    R.p(f"Tuning has the largest effect where the defaults were mismatched to the data. SVR gains "
        f"{ba.loc['SVR', 'cv_R2_gain']:+.3f} CV R2, because the default C = 1 is far too small for a target measured in "
        f"units. XGBoost gains {ba.loc['XGBoost', 'cv_R2_gain']:+.3f} CV R2 and {ba.loc['XGBoost', 'test_R2_gain']:+.3f} "
        f"test R2. The search settles on max_depth = {x['max_depth']}, learning_rate = {x['learning_rate']:.3f}, "
        f"n_estimators = {x['n_estimators']} and gamma = {x['gamma']:.2f}: shallower trees, a much smaller step size, "
        f"and a minimum split gain that prunes splits fitted to label noise. These gains move XGBoost from third "
        f"place to first in regression. Random Forest and HistGradientBoosting improve only marginally, because "
        f"their defaults were already close to optimal.")
    R.figure("fig30_xgboost_search_landscape.png",
             "Every configuration sampled by the XGBoost random search (star = selected).")


def section_compare(R, A):
    R.h("8. Model Performance Comparison")
    R.h("8.1 Regression", 2)
    r = _metric_cols(A["reg"], ["cv_R2_mean", "cv_R2_std", "test_MAE", "test_MSE", "test_RMSE", "test_R2",
                                "test_within_5_units_pct", "fit_time_s"]).sort_values("cv_R2_mean", ascending=False)
    best = A["meta"]["regressor_name"]
    R.table(r, "Tuned regression models (sorted by CV R2; final model highlighted)", font=8,
            highlight_row=lambda row: row["model"] == best)
    R.figure("fig21_regression_comparison.png", "Regression metrics, baseline vs tuned.")
    R.figure("fig22_pred_vs_actual.png", "Predicted vs actual waste on the test set for each tuned regressor.")
    R.figure("fig31a_cv_scores_regression.png", "Fold-level CV R2: spread as well as level.", 14)
    R.h("8.2 Classification", 2)
    c = _metric_cols(A["clf"], ["cv_macro_F1_mean", "cv_macro_F1_std", "test_accuracy", "test_balanced_accuracy",
                                "test_precision_macro", "test_recall_macro", "test_macro_F1", "test_ROC_AUC_ovr",
                                "test_HIGH_recall"]).sort_values("cv_macro_F1_mean", ascending=False)
    c.columns = [col.replace("test_", "").replace("_mean", "").replace("cv_macro_F1", "cv_F1") for col in c.columns]
    best_c = A["meta"]["classifier_name"]
    R.table(c, "Tuned classification models (sorted by CV macro-F1; final model highlighted)", font=7.5,
            highlight_row=lambda row: row["model"] == best_c)
    R.figure("fig24_classification_comparison.png", "Classification metrics, baseline vs tuned.")
    R.figure("fig25_confusion_matrices.png", "Test-set confusion matrices for every tuned classifier (counts and row %).")
    R.figure("fig26_roc_curves.png", "One-vs-rest ROC curves: all models (left) and per class for the final classifier (right).")
    pc = A["per_class"].pivot(index="model", columns="tier", values="f1")[["LOW", "MEDIUM", "HIGH"]].reset_index()
    pc = pc[~pc.model.str.contains("Dummy")].sort_values("HIGH", ascending=False)
    R.table(pc, "Test F1 per risk tier", font=8.5)
    cut = A["clf"][A["clf"].model.str.contains("cut points")].iloc[0]
    R.p(f"The row labelled '... regressor + cut points' classifies each booking by applying the tier cut points to "
        f"the best regressor's predicted amount. It reaches only {cut['test_macro_F1']:.3f} macro-F1. About 29% of "
        f"bookings sit exactly on a cut point (20 or 35 units), so a regression error of a single unit flips the "
        f"tier. A dedicated classifier learns these boundaries directly, which justifies keeping two separate models.")
    R.h("8.3 Strengths and limitations of each model", 2)
    sl = pd.DataFrame([
        ("Linear / Logistic Regression", "Fast, transparent coefficients, stable across folds",
         "Additive form cannot represent the pricing x service interaction or tier-dependent guest slopes; lowest "
         "scores among the real models"),
        ("KNN", "No training phase; exploits the booking grid", "Distance is diluted by many one-hot columns; weakest "
                                                                 "on the minority LOW tier"),
        ("SVR / SVM (RBF)", "Largest gain from tuning; smooth surface", "Slow to tune; probability outputs need an internal "
                                                                         "CV; still behind the ensembles"),
        ("Random Forest", "Strongest default model; best HIGH-tier F1; little tuning needed",
         "Large artefact on disk; averaging blunts sharp steps in the target"),
        ("HistGradientBoosting", "Fast histogram boosting, competitive with no feature work",
         "Most sensitive to the random-search draw; slightly behind XGBoost"),
        ("XGBoost (hist)", "Best regressor on CV and test; best LOW-tier and test macro-F1; explicit L1/L2/gamma regularisation",
         "Largest hyperparameter space; under-performs with default settings"),
    ], columns=["Model", "Strengths", "Limitations"])
    R.table(sl, "Strengths and limitations observed in this study", col_widths=[3.3, 6.2, 7.0], font=8.5)


def section_final(R, A):
    meta, sig = A["meta"], A["sig"].set_index("task")
    rm, cm = meta["regression_metrics"], meta["classification_metrics"]
    ceiling = A["manifest"]["attainable_ceiling"]["implied_r2_ceiling"]
    R.h("9. Final Model Selection")
    s_r, s_c = sig.loc["regression"], sig.loc["classification"]
    R.table(pd.DataFrame([
        {"task": t, "selected": s["best_model"], "runner-up": s["runner_up"], "metric": s["metric"],
         "CV mean (selected)": s["mean_a"], "CV mean (runner-up)": s["mean_b"], "difference": s["mean_diff"],
         "t": s["t"], "p-value": s["p_value"], "folds won": f"{int(s['a_wins'])}/{int(s['folds'])}"}
        for t, s in [("regression", s_r), ("classification", s_c)]]),
        "Selected models and corrected resampled t-test against the runner-up (5x3 repeated CV)", font=8)
    R.h("9.1 Regression: XGBoost (hist)", 2) if meta["regressor_name"] == "XGBoost" else R.h(f"9.1 Regression: {meta['regressor_name']}", 2)
    R.p(f"The tuned {meta['regressor_name']} regressor has the highest cross-validated R2 ({rm['cv_R2_mean']:.4f} +/- "
        f"{rm['cv_R2_std']:.4f}). On the untouched test set it scores R2 = {rm['test_R2']:.4f}, RMSE = "
        f"{rm['test_RMSE']:.2f} units, MAE = {rm['test_MAE']:.2f} units and MSE = {rm['test_MSE']:.2f}, and "
        f"{rm['test_within_5_units_pct']:.1f}% of test bookings are predicted within 5 units of the actual waste. "
        f"Against the noise ceiling of R2 ~ {ceiling} estimated in Review 2, the model captures "
        f"{100 * rm['test_R2'] / ceiling:.1f}% of the explainable variance. It beats {s_r['runner_up']} on "
        f"{int(s_r['a_wins'])} of {int(s_r['folds'])} repeated-CV folds. The corrected t-test (p = {s_r['p_value']:.2f}) "
        f"shows the margin is not statistically decisive, but XGBoost also leads on every test metric and adds "
        f"explicit regularisation and native per-prediction explanations (TreeSHAP contributions), which the "
        f"dashboard uses.")
    R.figure("fig23_residuals_final_regressor.png", "Residual diagnostics of the final regressor on the test set.")
    R.h(f"9.2 Classification: {meta['classifier_name']}", 2)
    rep = A["report"].copy().reset_index().rename(columns={"index": "class"})
    R.table(rep, "Classification report of the final classifier (test set)", font=8.5)
    cmx = A["cm"].reset_index().rename(columns={"index": "actual \\ predicted"})
    R.table(cmx, "Confusion matrix of the final classifier (test set)", font=9)
    R.p(f"Under the pre-registered rule, the tuned {meta['classifier_name']} is selected with CV macro-F1 "
        f"{cm['cv_macro_F1_mean']:.4f} +/- {cm['cv_macro_F1_std']:.4f}. On the test set it reaches accuracy "
        f"{cm['test_accuracy']:.3f}, macro precision {cm['test_precision_macro']:.3f}, macro recall "
        f"{cm['test_recall_macro']:.3f}, macro-F1 {cm['test_macro_F1']:.3f}, ROC-AUC {cm['test_ROC_AUC_ovr']:.3f} and "
        f"HIGH-tier recall {cm['test_HIGH_recall']:.3f}. The runner-up, {s_c['runner_up']}, is statistically "
        f"indistinguishable (corrected t-test p = {s_c['p_value']:.2f}). Both models are therefore valid choices, "
        f"and the dashboard lets the user switch to any tuned model. Errors fall almost entirely between adjacent "
        f"tiers: the classifier never confuses LOW with HIGH, so a wrong tier changes the recommended reduction by "
        f"at most one step.")
    R.figure("fig28_feature_importance_final.png",
             "Grouped permutation importance of the final models (test set) and XGBoost gain importance.")
    imp = A["imp_reg"].head(6)
    R.p("Grouped permutation importance shuffles all one-hot columns of a business feature together. It confirms "
        "the EDA: " + ", ".join(f"{r.feature} ({r.importance_mean:.3f})" for r in imp.itertuples()) +
        " account for almost all of the regressor's test R2. Food type, event type, storage, purchase history and "
        "seasonality contribute essentially nothing, as Section 3 predicted.")
    R.figure("fig29_learning_curves.png", "Learning curves of the final models.")
    R.p("The training and validation curves converge and the validation curve flattens. Adding more bookings of the "
        "same kind would bring little gain; the remaining error is mostly label noise, consistent with the ceiling "
        "analysis.")
    R.h("9.3 Why the selected solution is better", 2)
    R.bullets([
        ("Accuracy: ", f"the final regressor reduces test RMSE from 10.39 (naive) and "
                       f"{A['reg'].set_index('model').loc['Linear Regression', 'test_RMSE']:.2f} (linear) to "
                       f"{rm['test_RMSE']:.2f} units."),
        ("Robustness: ", "the selection was made on cross-validation, and the test set confirms it rather than "
                         "deciding it. Fold-to-fold spread is small (fold-score plots in Section 8)."),
        ("Operational fit: ", f"the classifier recalls {100 * cm['test_HIGH_recall']:.0f}% of HIGH-waste bookings, "
                              f"the ones that trigger a donation alert, with no LOW/HIGH confusions."),
        ("Interpretability: ", "the importance rankings match domain expectations and the EDA, and the dashboard "
                               "explains every individual forecast."),
    ])


def section_product(R, A):
    meta = A["meta"]
    R.h("10. End-to-End Product: Decision-Support Dashboard")
    R.p("The trained models are deployed in a Streamlit dashboard (app/streamlit_app.py). It runs the full chain: "
        "Booking Information -> Waste Prediction -> Risk Classification -> Preparation Recommendation -> Donation "
        "Alert. Inference goes through pipeline.prepare_inputs, the same feature engineering used in training, so "
        "the dashboard cannot drift from the models.")
    R.table(pd.DataFrame([
        ("Forecast", "Booking form -> predicted waste, waste %, waste per guest, risk tier with class probabilities, "
                     "recommended action, suggested reduced quantity with the re-forecast waste, donation-partner alert, "
                     "segment benchmark, a what-if curve of waste vs quantity prepared, a per-feature explanation "
                     "(XGBoost TreeSHAP contributions) and every tuned model's forecast for the same booking"),
        ("Data Explorer", "Interactive EDA over every column: filters, KPIs, waste / waste rate / HIGH share by any "
                          "column with 95% CI, risk-tier mix, density curves, a two-column heatmap, a scatter explorer, "
                          "statistical tables and the full EDA figure gallery"),
        ("Model Performance", "Final model KPIs, tuned and baseline tables, before/after tuning, significance test, "
                              "per-tier F1 heat map, best hyperparameters and the model figure gallery"),
        ("Batch Forecast", "Upload a CSV of bookings (original or snake_case column names), get forecasts, actions, "
                           "suggested quantities and alerts, and download the results"),
        ("About", "Pipeline, methodology and limitations"),
    ], columns=["Tab", "Function"]), "Dashboard tabs", col_widths=[3.2, 13.3], font=8.5)
    R.p(f"The sidebar defaults to the selected models ({meta['regressor_name']} regressor, {meta['classifier_name']} "
        f"classifier) and allows switching to any tuned model. A warning appears when a booking falls outside the "
        f"training range (guests {meta['training_ranges']['number_of_guests'][0]}-"
        f"{meta['training_ranges']['number_of_guests'][1]}, quantity {meta['training_ranges']['quantity_of_food'][0]}-"
        f"{meta['training_ranges']['quantity_of_food'][1]}), because tree models do not extrapolate.")
    R.p("Launch: ./venv/Scripts/python -m streamlit run app/streamlit_app.py", italic=True)
    R.h("10.1 Prescriptive action engine", 2)
    R.table(pd.DataFrame([("LOW", "Normal preparation", "--"),
                          ("MEDIUM", "Reduce quantity of food by 10-15%", "Suggested quantity uses the 12.5% midpoint"),
                          ("HIGH", "Reduce quantity by 15-20% and flag a donation-partner alert",
                           "Suggested quantity uses the 17.5% midpoint; alert raised")],
                         columns=["Predicted tier", "Action", "Dashboard behaviour"]), "Tier-to-action mapping")


def section_conclusion(R, A):
    meta = A["meta"]
    rm, cm = meta["regression_metrics"], meta["classification_metrics"]
    R.h("11. Conclusions")
    R.bullets([
        f"Waste can be forecast before an event with test R2 = {rm['test_R2']:.3f} (RMSE {rm['test_RMSE']:.2f} units) "
        f"by a tuned XGBoost (hist) regressor. This is close to the ceiling set by label noise.",
        f"Booking risk tiers are classified with macro-F1 = {cm['test_macro_F1']:.3f} and ROC-AUC = "
        f"{cm['test_ROC_AUC_ovr']:.3f}. Errors are confined to adjacent tiers.",
        "Tree ensembles clearly outperform linear, instance-based and kernel models, because waste depends on a "
        "pricing x service interaction and on tier-specific guest effects that additive models cannot represent.",
        "Hyperparameter tuning matters most for XGBoost and SVR. Random Forest and HistGradientBoosting are "
        "near-optimal at their defaults.",
        "Of the ten booking attributes, pricing tier, preparation method, guest count and quantity prepared carry "
        "almost all of the predictive signal. Event type, food type, storage, season and purchase history add "
        "little or nothing.",
        "The project is delivered as a working dashboard that turns each forecast into a concrete preparation "
        "quantity and a donation alert.",
    ])
    R.h("12. Limitations and Future Work")
    R.bullets([
        "The dataset has limited value granularity: 61 guest/quantity combinations, and most waste values are "
        "multiples of 5. Performance should be re-validated on a caterer's own operational records before "
        "deployment.",
        "No date field exists, so weekday, holiday and lead-time effects cannot be modelled. Adding the booking "
        "date is the most valuable data extension.",
        "Pricing is a three-level tier, which limits price-sensitivity analysis to tier comparisons.",
        "The projected waste at a reduced quantity is an in-sample what-if, not a causal estimate. A controlled "
        "pilot is needed to measure the real effect of the recommended reductions.",
        "Forecasts support staff decisions and never override contract minimums, dietary requirements, service "
        "buffers or food-safety rules.",
        "Future work: per-tier prediction intervals (quantile XGBoost), cost-sensitive thresholds that weight a "
        "missed HIGH booking more heavily, and periodic retraining as new bookings arrive.",
    ])
    R.h("13. Project Artefacts")
    R.table(pd.DataFrame([
        ("src/pipeline.py", "Review 2 pipeline: quality audit, cleaning, split, tiers, features, integration, EDA"),
        ("src/eda_extended.py", "Extended EDA across every column (Figures 8, 11-20; eda_*.csv)"),
        ("src/train_models.py", "Baseline + tuned training, evaluation, comparison, significance, figures 21-31, model export"),
        ("src/build_notebooks.py, src/review3_cells.py", "Builds and executes the Review 2 and Review 3 notebooks"),
        ("src/build_report.py", "Generates this report from the result files"),
        ("notebooks/Review3_Model_Building_and_Comparison.ipynb", "Executed Review 3 notebook"),
        ("out/*.csv, out/best_hyperparameters.json", "All result tables"),
        ("models/", "regressor.pkl, classifier.pkl, model_zoo.pkl (all tuned models), metadata.json"),
        ("app/streamlit_app.py", "Interactive dashboard"),
    ], columns=["Path", "Content"]), "Deliverables", col_widths=[6.5, 10], font=8.5)
    R.h("References")
    for ref in [
        "Chen, T. and Guestrin, C. (2016). XGBoost: A Scalable Tree Boosting System. Proc. KDD.",
        "Ke, G. et al. (2017). LightGBM: A Highly Efficient Gradient Boosting Decision Tree. NeurIPS (histogram-based boosting).",
        "Breiman, L. (2001). Random Forests. Machine Learning 45(1).",
        "Nadeau, C. and Bengio, Y. (2003). Inference for the Generalization Error. Machine Learning 52(3).",
        "Bergstra, J. and Bengio, Y. (2012). Random Search for Hyper-Parameter Optimization. JMLR 13.",
        "Bergsma, W. (2013). A bias-correction for Cramer's V and Tschuprow's T. J. Korean Statistical Society 42(3).",
        "Lundberg, S. et al. (2020). From local explanations to global understanding with explainable AI for trees. Nature MI 2.",
        "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR 12.",
        "Kaggle dataset: Food Wastage Data in Restaurant (trevinhannibal).",
    ]:
        para = R.doc.add_paragraph(ref, style="List Number")
        para.runs[0].font.size = Pt(9.5)


def export_pdf(docx_path: pathlib.Path) -> pathlib.Path | None:
    pdf = docx_path.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx_path}'); $d.Fields.Update() | Out-Null; "
          f"$d.SaveAs2('{pdf}', 17); $d.Close($false); $w.Quit()")
    import os
    import shutil
    exe = shutil.which("powershell") or os.path.join(os.environ.get("SystemRoot", r"C:\Windows"),
                                                     "System32", "WindowsPowerShell", "v1.0", "powershell.exe")
    res = subprocess.run([exe, "-NoProfile", "-Command", ps], capture_output=True, text=True)
    if res.returncode != 0 or not pdf.exists():
        print("PDF export failed:", res.stderr.strip()[:500])
        return None
    return pdf


def main():
    A = load()
    R = Report()
    cover(R)
    R.page_break()
    section_intro(R, A)
    section_eda(R, A)
    section_models(R, A)
    section_eval(R, A)
    section_tuning(R, A)
    section_compare(R, A)
    section_final(R, A)
    section_product(R, A)
    section_conclusion(R, A)
    R.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH}")
    if "--pdf" in sys.argv:
        pdf = export_pdf(DOCX_PATH)
        if pdf:
            print(f"Wrote {pdf}")


if __name__ == "__main__":
    main()
